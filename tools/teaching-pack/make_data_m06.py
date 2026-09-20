"""Deterministic M06 traffic evidence and spoiler-safe gauge workbook."""
import csv
import os
from collections import defaultdict
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(ROOT, "..", ".."))
M05 = os.path.join(REPO, "teaching-pack", "09-M05-overnight-alarm")
OUT = os.path.join(ROOT, "build", "M06")
os.makedirs(OUT, exist_ok=True)

feed_path = os.environ.get("M05_FEED", os.path.join(M05, "feed.csv"))
if not os.path.exists(feed_path):
    raise FileNotFoundError(
        "M06 depends on M05 feed.csv from PR #9. Merge/check out that dependency, "
        "then rerun tools/teaching-pack/build_m06.sh."
    )

traffic_path = os.path.join(OUT, "traffic_counts.csv")
start = datetime(2027, 2, 16, 2, 30)
traffic = []
classes = ["2-axle passenger", "2-axle passenger", "2-axle light truck"]
for i in range(13):
    ts = start + timedelta(minutes=5 * i)
    vehicle_class = classes[i % len(classes)]
    axles = 2
    weight = 3400 + ((i * 1300) % 3900)
    direction = "EB" if i % 2 else "WB"
    if ts == datetime(2027, 2, 16, 3, 0):
        vehicle_class, axles, weight, direction = "5-axle heavy vehicle", 5, 79400, "EB"
    traffic.append({
        "timestamp_cst": ts.strftime("%Y-%m-%d %H:%M"),
        "direction": direction,
        "vehicle_class": vehicle_class,
        "axle_count": axles,
        "estimated_gross_weight_lb": weight,
        "source": "Kinnick County bridge traffic counter",
        "quality_flag": "OK",
    })

with open(traffic_path, "w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(traffic[0]))
    writer.writeheader()
    writer.writerows(traffic)

with open(feed_path, newline="") as f:
    feed = list(csv.DictReader(f))

# Preserve the duplicate callback in the source file, but do not let duplicated payloads
# silently double-count the daily peak table.
unique = {}
for row in feed:
    unique[(row["timestamp_cst"], row["gauge_id"])] = row

daily = defaultdict(lambda: {"value": float("-inf"), "timestamp": "", "gauge": ""})
for (timestamp, gauge), row in unique.items():
    value = float(row["corrected_strain_microstrain"])
    date = timestamp[:10]
    if value > daily[date]["value"]:
        daily[date] = {"value": value, "timestamp": timestamp, "gauge": gauge}

wb = Workbook()
ws = wb.active
ws.title = "Daily peaks"
headers = [
    "date", "peak_timestamp_cst", "gauge_id", "daily_peak_corrected_microstrain",
    "keep_or_question", "evidence_note",
]
ws.append(headers)
for date in sorted(daily):
    row = daily[date]
    ws.append([date, row["timestamp"], row["gauge"], row["value"], "", ""])

event = wb.create_sheet("03-00 event")
event.append([
    "timestamp_cst", "gauge_id", "corrected_strain_microstrain",
    "traffic_record_checked", "interpretation",
])
for gauge in [f"G{i}" for i in range(1, 9)]:
    row = unique[("2027-02-16 03:00", gauge)]
    event.append([row["timestamp_cst"], gauge, float(row["corrected_strain_microstrain"]), "", ""])

decision = wb.create_sheet("Decision log")
decision.append(["item", "action chosen", "project hours", "evidence checked", "accepted unverified"])
decision.append(["M06 alarm audit", "", "", "", ""])

handoff = wb.create_sheet("Wes handoff")
handoff.append(["field", "student entry"])
for label in [
    "Student / file owner", "Filename", "Date handed to Wes", "What I checked",
    "What I did not check", "Known limitations", "Decision-log row included?",
]:
    handoff.append([label, ""])
handoff.append(["RETAIN THIS RECORD", "This February file returns later in the project."])

blue = "1F3A5F"
orange = "F26B1D"
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column in sheet.columns:
        width = min(42, max(12, max(len(str(c.value or "")) for c in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1

handoff["A10"].font = Font(bold=True, color=orange)
wb.save(os.path.join(OUT, "M06-gauge-audit-and-handoff.xlsx"))

print(f"  wrote M06/traffic_counts.csv ({len(traffic)} rows)")
print(f"  wrote M06/M06-gauge-audit-and-handoff.xlsx ({len(daily)} daily peaks)")
