"""M04 decks — spoiler-safe descriptive-statistics hunt and instructor reveal."""
import json
import os
import sys
from io import BytesIO

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from slidelib import *

C = json.load(open(os.path.join(ROOT, "build", "facts.json")))["m04"]


def add_speaker_notes(prs, notes):
    if len(prs.slides) != len(notes):
        raise ValueError(f"M04 notes count ({len(notes)}) != slides ({len(prs.slides)})")
    for slide, note in zip(prs.slides, notes):
        slide.notes_slide.notes_text_frame.text = note.strip()


def slides_from(source, numbers):
    requested = list(numbers)
    if len(source.slides) != 20:
        raise ValueError(f"M04 source deck must have 20 slides, found {len(source.slides)}")
    buffer = BytesIO()
    source.save(buffer)
    buffer.seek(0)
    result = Presentation(buffer)
    keep = set(requested)
    for index in reversed(range(len(result.slides))):
        if index + 1 not in keep:
            slide_id = result.slides._sldIdLst[index]
            result.part.drop_rel(slide_id.rId)
            del result.slides._sldIdLst[index]
    return result


prs = deck()

# 1
title_slide(prs, "Meeting 04  ·  Unit 1: Hired",
            "One number.",
            "Descriptive statistics, and the information a center cannot carry",
            "Thursday, February 11, 2027  ·  SEIS 201")

# 2
s, y = content(prs, "The laboratory sent 24 deck-core results",
               "Diane has a county briefing and one line of space")
speaker_card(s, 0.72, y + 0.15, 11.87, 2.05, "DIANE HALVORSEN, PE",
             ["“Give me one number for the deck. I can explain one number.”"],
             body_size=24)
card(s, 0.72, y + 2.5, 11.87, 1.6, "THE ASK",
     ["One number for center. One number for spread. A count below the 4,500 psi project "
      "comparison. Then three sentences she can repeat."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=17)
text(s, 0.72, y + 4.35, 11.9, 0.55,
     "Six project hours. You choose which parts of the answer to verify.",
     size=17, color=GREY, italic=True)

# 3
s, y = content(prs, "FOREMAN already gave her the one number")
speaker_card(s, 0.72, y + 0.15, 11.87, 2.25,
             "FOREMAN v4.2 — DECK CORE EXECUTIVE SUMMARY",
             [f"“ADEQUATE. Mean core strength is {C['mean_psi']:,.0f} psi, exceeding the "
              f"{C['comparison_psi']:,} psi project value by "
              f"{C['mean_psi'] - C['comparison_psi']:,.0f} psi. No further statistical "
              "review is required. Confidence 94%.”"],
             machine=True, body_size=19)
card(s, 0.72, y + 2.75, 11.87, 1.5, "YOUR QUESTION",
     ["Is that a good summary of the file? Do not answer from the paragraph. Answer from the "
      "24 rows."], tint=PAPER, edge=CONCRETE, label_color=GIRDER, body_size=17)

# 4
statement(prs, "One number can be correct\nand still be a bad summary.",
          "Today the hunt is to identify what the compression throws away.")

# 5
section(prs, "1", "Build the summaries",
        "Center, spread, and scale — three different jobs")

# 6
s, y = content(prs, "First: name what is in the file",
               "A dataset is not just a rectangle of numbers")
bullet_list(s, 0.95, y + 0.3, 11.4, 3.5, [
    ("Observational unit  ", "the thing represented by one row. Here, one tested core."),
    ("Variable  ", "a recorded feature that can change from row to row."),
    ("Quantitative variable  ", "a measured amount on which arithmetic is meaningful."),
    ("Categorical variable  ", "a label such as side or drill zone; it groups observations."),
    ("Dataset  ", "the observations and variables together, plus enough context to interpret them."),
], size=18, dot=GIRDER_LT)
card(s, 0.72, y + 3.95, 11.87, 1.35, "BEFORE ANY FORMULA",
     ["Say what one row is and which variable you are summarizing."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=14)

# 7
s, y = content(prs, "Mean — where the values balance",
               "Useful, reproducible, and blind to where individual values sit")
rect(s, 0.72, y + 0.25, 11.87, 1.25, fill=PAPER, line=CONCRETE, radius=0.05)
text(s, 0.72, y + 0.52, 11.87, 0.7, "x̄  =  Σ xᵢ / n", size=37, color=GIRDER,
     font="Courier New", bold=True, align=PP_ALIGN.CENTER)
bullet_list(s, 0.95, y + 1.95, 11.4, 2.0, [
    ("It answers: ", "where is the center of these values?"),
    ("It does not answer: ", "how far apart are they, how many are low, or where the sample came from."),
    ("Units stay the same: ", "psi in, psi out."),
], size=19, dot=STICKER)
card(s, 0.72, y + 4.13, 11.87, 1.30, "SPREADSHEET",
     ["=AVERAGE(range)"], tint=PAPER, edge=CONCRETE, label_color=GIRDER,
     body_size=14, mono=True)

# 8
s, y = content(prs, "Variance and standard deviation",
               "How far the observations typically sit from the mean")
card(s, 0.72, y + 0.2, 5.85, 3.2, "SAMPLE VARIANCE  ·  s²",
     ["Square each distance from the mean, add them, divide by n − 1.",
      "",
      "Units are squared: psi².",
      "",
      "=VAR.S(range)"], tint=PAPER, edge=CONCRETE, label_color=GIRDER,
     body_size=17)
card(s, 6.72, y + 0.2, 5.87, 3.2, "SAMPLE STANDARD DEVIATION  ·  s",
     ["Take the square root of variance.",
      "",
      "Back in the original units: psi.",
      "",
      "=STDEV.S(range)"], tint=CREAM, edge=GOLD, label_color=REBAR,
     body_size=17)
text(s, 0.95, y + 3.75, 11.4, 1.0,
     "Same mean + different standard deviation = different engineering story.",
     size=21, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 9
s, y = content(prs, "Coefficient of variation puts spread on a relative scale",
               "Useful when the size of the mean changes")
rect(s, 0.72, y + 0.25, 11.87, 1.25, fill=PEACH, line=STICKER, radius=0.05)
text(s, 0.72, y + 0.5, 11.87, 0.72, "CV  =  s / x̄  ×  100%", size=36,
     color=STICKER, font="Courier New", bold=True, align=PP_ALIGN.CENTER)
bullet_list(s, 0.95, y + 1.95, 11.4, 2.2, [
    ("Standard deviation  ", "keeps the engineering units. It says how many psi."),
    ("Coefficient of variation  ", "removes the units. It says how large the spread is relative to the mean."),
    ("Read it as a comparison, not a verdict. ", "CV has no universal pass/fail boundary."),
], size=18, dot=STICKER)
card(s, 0.72, y + 4.13, 11.87, 1.30, "SPREADSHEET",
     ["=STDEV.S(range)/AVERAGE(range)  ·  format as percent"],
     tint=PAPER, edge=CONCRETE, label_color=GIRDER, body_size=14, mono=True)

# 10
s, y = content(prs, "Sample or population?",
               "The formula is a claim about what your file represents")
speaker_card(s, 0.72, y + 0.2, 5.85, 3.25, "SAMPLE  ·  USE .S",
             ["The 24 cores are evidence about a larger physical deck.",
              "",
              "Use STDEV.S and VAR.S.",
              "",
              "The n − 1 denominator accounts for estimating spread from a sample."],
             body_size=15)
card(s, 6.72, y + 0.2, 5.87, 3.25, "POPULATION  ·  USE .P",
     ["Use only when the file contains every member of the population you intend to describe.",
      "",
      "STDEV.P answers a different question. It is not the “more exact” button."],
     tint=PAPER, edge=CONCRETE, label_color=GIRDER, body_size=17)
card(s, 0.72, y + 3.75, 11.87, 1.35, "TODAY",
     ["Twenty-four tested cores are the full file, but they are not the whole deck."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=15)

# 11
activity(prs, "Hunt the story FOREMAN compressed", 22, [
    "Open cores_2027.csv. Name the observational unit, quantitative variable, and location variables.",
    "Compute n, mean, sample variance, sample SD, CV, minimum, maximum, and the count below 4,500 psi.",
    "Write one sentence using only the mean. Then write one that adds spread and the low-result count.",
    "Read the location columns. State what these 24 rows can and cannot say about the whole deck.",
], note="Keep the result private until every pair has center, spread, and count. H4-02 carries the formulas.")

# 12
statement(prs, "Find two sentences that are\nboth numerically correct.",
          "Then decide which one keeps the decision-relevant information.", dark=False)

# 13
image_slide(prs, "Put every core back on the page", "m04-core-spread.png",
            kicker="The orange line is the mean. Every dot is still an observation.",
            caption="The mean sits above the project comparison while individual results extend "
                    "well below it. Center and spread are answering different questions.",
            box_h=4.15)

# 14
rows = [
    ["Statistic", "Result", "Job"],
    ["n", str(C["n"]), "How much data"],
    ["Mean", f"{C['mean_psi']:,.2f} psi", "Center"],
    ["Sample variance", f"{C['variance_psi2']:,.2f} psi²", "Squared spread"],
    ["Sample SD", f"{C['sample_sd_psi']:,.2f} psi", "Spread in original units"],
    ["Sample CV", f"{C['sample_cv_pct']:.2f}%", "Spread relative to center"],
    ["Range", f"{C['minimum_psi']:,}–{C['maximum_psi']:,} psi", "Extremes"],
    ["Below 4,500", f"{C['below_comparison']} of {C['n']}", "Low tail"],
]
table_slide(prs, "One file, several honest summaries", rows,
            widths=[2.5, 3.0, 4.4], size=14,
            col_align=["l", "r", "l"], row_colors={7: CREAM})

# 15
s, y = content(prs, "The mean did not move. The story did.")
speaker_card(s, 0.72, y + 0.2, 5.85, 3.45, "MEAN-ONLY SENTENCE",
             [f"“The 24 cores average {C['mean_psi']:,.0f} psi, above the "
              f"{C['comparison_psi']:,} psi project comparison.”"],
             body_size=20)
card(s, 6.72, y + 0.2, 5.87, 3.45, "SPREAD-AWARE SENTENCE",
     [f"“The mean is {C['mean_psi']:,.0f} psi, but sample SD is "
      f"{C['sample_sd_psi']:,.0f} psi (CV {C['sample_cv_pct']:.1f}%) and "
      f"{C['below_comparison']} of {C['n']} results are below 4,500 psi.”"],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=19)
text(s, 0.72, y + 4.0, 11.9, 0.85,
     "Both sentences are numerically correct. Only one tells the reader where the risk in the "
     "summary lives.", size=19, color=REBAR, align=PP_ALIGN.CENTER)

# 16
s, y = content(prs, "Uncertainty from spread",
               "The file does not support a single-value picture of the tested concrete")
bullet_list(s, 0.95, y + 0.35, 11.4, 3.1, [
    ("Spread is not a mistake to clean away. ", "It is observed variation in the sampled material and testing process."),
    ("Standard deviation quantifies it. ", "CV makes its size comparable to the mean."),
    ("The low tail remains real after averaging. ", "A center summarizes observations; it does not replace them."),
    ("This is descriptive uncertainty. ", "Inference about the whole deck requires attention to how and where the sample was taken."),
], size=18, dot=GIRDER_LT)
card(s, 0.72, y + 3.78, 11.87, 1.55, "DO NOT OVERCLAIM",
     ["A core below 4,500 psi is not, by itself, a complete code finding. A mean above 4,500 psi "
      "is not, by itself, an adequacy determination."],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=15)

# 17
s, y = content(prs, "Twenty-four cores are not the deck",
               "Read the columns FOREMAN did not summarize")
card(s, 0.72, y + 0.2, 11.87, 1.55, "WHAT THE FILE SAYS",
     ["Every row lists drill_zone = shoulder. Access was easy there. No core in this file came "
      "from a traffic lane, joint, tower interface, or underside."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=15)
bullet_list(s, 0.95, y + 2.15, 11.4, 2.25, [
    ("Population  ", "every location or volume of deck concrete about which the claim is intended."),
    ("Sample  ", "the 24 tested cores."),
    ("Scope  ", "the results describe these sampled shoulder locations. They do not establish the condition of every deck region."),
], size=18, dot=GIRDER_LT)
text(s, 0.72, y + 4.55, 11.9, 0.6,
     "More rows do not repair a sampling frame that never reaches the places you care about.",
     size=18, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 18
s, y = content(prs, "A three-sentence answer Diane can use")
card(s, 0.72, y + 0.2, 11.87, 3.75, "MODEL — REVEAL AFTER STUDENTS WRITE",
     [f"1. The 24 tested cores average {C['mean_psi']:,.0f} psi.",
      f"2. Their sample SD is {C['sample_sd_psi']:,.0f} psi "
      f"(CV {C['sample_cv_pct']:.1f}%), and {C['below_comparison']} of {C['n']} results fall "
      "below the 4,500 psi project comparison, so the mean alone hides consequential spread.",
      "3. Because every core was drilled from an accessible shoulder location, these results "
      "describe the sampled locations but do not establish the condition of the entire deck."],
     tint=RGBColor(0xEE, 0xF3, 0xF8), edge=GIRDER_LT, label_color=GIRDER,
     body_size=18)
text(s, 0.95, y + 4.25, 11.4, 0.65,
     "Short does not have to mean false. It has to mean disciplined.",
     size=20, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 19
s = blank(prs)
bg(s, prs, WHITE)
head(s, "Mask off — two minutes")
bullet_list(s, 0.95, 1.9, 11.4, 3.4, [
    ("Diane is wrong to demand only one number. ", "The mean cannot carry the spread, the low tail, and the sampling scope."),
    ("Diane is right to demand a short answer. ", "Clients need a usable statement, not a pasted spreadsheet."),
    ("FOREMAN's arithmetic is correct. ", "The designed error is turning one correct statistic into an unsupported adequacy conclusion."),
    ("Diane is not the grader. ", "The check and the record are graded, not whether she likes the brief."),
], size=18, dot=GIRDER_LT)

# 20
closer(prs, [
    ("Hand in", "H4-04: three sentences plus the one-sentence verification note"),
    ("Log", "What you checked, what you accepted, and how many of the six hours you spent"),
    ("Tuesday", "FOREMAN runs against the live sensor stream overnight — and acts on one reading"),
], title="Before Tuesday")

student_slides = list(range(1, 13))
reveal_slides = list(range(13, 21))
if set(student_slides) & set(reveal_slides) or set(student_slides + reveal_slides) != set(range(1, 21)):
    raise ValueError("M04 student and reveal decks must partition all slides")

student = slides_from(prs, student_slides)
add_speaker_notes(student, [
    "Open with the job, not the formulas. Students will receive the numerical problem before the vocabulary.",
    "Read Diane's line without mocking it. She needs brevity; the lesson is what disciplined brevity requires.",
    "Take a vote: ready to brief, not ready, or unsure. Do not validate a choice and do not disclose any omitted statistic.",
    "Emphasize that a correct calculation can still be an incomplete summary. Do not name which summary is incomplete yet.",
    "Transition to the six Week 2 concepts. Students need all three jobs—center, spread, scale—before opening the file.",
    "Ask students to name one row and one quantitative variable from a familiar CSV before using the core file.",
    "Keep mean concrete: balance point and same units. Ask what information has vanished when 24 rows become one value.",
    "Explain why variance has squared units and why SD returns to psi. Do not compute the M04 values on screen.",
    "CV is dimensionless relative spread. Explicitly reject a universal CV pass/fail cutoff.",
    "The file is complete as a file but incomplete as the deck. That distinction is the sample/population concept.",
    "Start the 22-minute hunt. Require each pair to have center, spread, count, and scope before discussion.",
    "Leave this slide up while pairs prepare two numerically correct sentences. Do not switch decks until students articulate the contrast.",
])
save(student, os.path.join(ROOT, "build", "M04", "M04-student.pptx"))

reveal = slides_from(prs, reveal_slides)
add_speaker_notes(reveal, [
    "Open only after pairs report their values. Ask students to narrate the plot before you narrate it.",
    "Reconcile spreadsheet outputs. STDEV.S and VAR.S are the expected formulas because these cores sample a larger deck.",
    "Read both sentences aloud. Ask what disappeared in the mean-only version and what decision might change when it returns.",
    "Frame spread as a finding, not a defect in the dataset. Hold the engineering boundary: descriptive statistics are not a code determination.",
    "Reveal drill_zone = shoulder. Ask which deck regions remain unobserved and why 24 convenient cores are still a sample.",
    "Show the model only after students write. Accept equivalent language that gives center, spread/count, and scope without overclaiming.",
    "Mask off the role tension: Diane's brevity request is operationally reasonable, while one-number certainty is not.",
    "Collect H4-04 and decision-log entries. Preview M05 without revealing its designed error.",
])
save(reveal, os.path.join(ROOT, "build", "M04", "M04-instructor-reveal.pptx"))
