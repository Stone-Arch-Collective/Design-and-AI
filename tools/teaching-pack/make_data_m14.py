"""Deterministic M14 Exam 1 evidence for the overweight-permit scenario."""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "build", "M14")
os.makedirs(OUT, exist_ok=True)

TEST_LOAD_KIP = 80.0
PERMIT_LOAD_KIP = 104.0
SCREENING_TRIGGER = 240.0
T_95_DF7 = 2.365
STRAINS = [156.0, 162.0, 168.0, 171.0, 176.0, 181.0, 185.0, 190.0]


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


wb = Workbook()
start = wb.active
start.title = "START_HERE"
start.append(["item", "value"])
start_rows = [
    ["release_status", "READY FOR SHANNON SKIM · NOT LMS PUBLISHED"],
    ["exam", "SEIS 201 Exam 1 · M14"],
    ["job", "Otter Bend Lift Bridge · ACMEJOB Unit 2"],
    ["decision", "Internal screening recommendation for one requested night crossing"],
    ["time limit", "75 minutes"],
    ["test load", f"{TEST_LOAD_KIP:.0f} kip"],
    ["requested permit load", f"{PERMIT_LOAD_KIP:.0f} kip"],
    ["instructional screening trigger", f"{SCREENING_TRIGGER:.0f} microstrain"],
    ["scope", "Use only the supplied clean exam file and FOREMAN recommendation."],
    ["professional boundary", "This fictional trigger is not a legal permit criterion or bridge capacity rating."],
]
for row in start_rows:
    start.append(row)
style_sheet(start, [34, 84])

permit = wb.create_sheet("PERMIT_REQUEST")
permit.append(["source_field", "source_value", "source_note"])
permit_rows = [
    ["request_id", "KC-OW-2027-0318", "Kinnick County Public Works"],
    ["crossing", "one crossing", "No return crossing is included"],
    ["requested_time", "2027-03-18 23:30 CST", "Night crossing"],
    ["gross_vehicle_weight_kip", f"{PERMIT_LOAD_KIP:.1f}", "Requested gross crossing load"],
    ["maximum_speed_mph", "10", "Requested operating condition"],
    ["lane_position", "center lane", "Requested operating condition"],
    ["test_truck_weight_kip", f"{TEST_LOAD_KIP:.1f}", "Reference load used for supplied control-gauge test"],
]
for row in permit_rows:
    permit.append(row)
style_sheet(permit, [33, 35, 55])

test = wb.create_sheet("CONTROL_GAUGE_TEST")
test.append([
    "pass_id", "test_truck_weight_kip", "gauge_id",
    "peak_strain_microstrain", "speed_mph", "lane_position", "quality_flag",
])
for idx, strain in enumerate(STRAINS, 1):
    test.append([
        f"LT-{idx:02d}", f"{TEST_LOAD_KIP:.1f}", "G4",
        f"{strain:.1f}", "10", "center", "VALID",
    ])
style_sheet(test, [14, 24, 13, 27, 14, 18, 16])

criteria = wb.create_sheet("REVIEW_BASIS")
criteria.append(["basis_item", "value", "interpretation"])
criteria_rows = [
    ["instructional screening trigger", f"{SCREENING_TRIGGER:.1f} microstrain",
     "A projected interval reaching this trigger requires escalation; this is not a capacity limit."],
    ["scaling method for this exam", "permit load / test load",
     "Apply a direct ratio only as the stated classroom screening simplification."],
    ["95% interval multiplier", f"t* = {T_95_DF7:.3f}",
     "Use for n = 8 repeated valid passes; df = 7."],
    ["known limitation", "linear response is assumed, not verified",
     "Carry this limitation into the recommendation."],
    ["known limitation", "one control gauge and one test condition",
     "Do not claim a bridge-wide capacity determination."],
]
for row in criteria_rows:
    criteria.append(row)
style_sheet(criteria, [34, 36, 68])

wb.properties.title = "SEIS 201 M14 Exam 1 clean evidence file"
wb.properties.subject = "Otter Bend overweight-permit applied practical"
wb.properties.creator = "SEIS 201"
wb.save(os.path.join(OUT, "M14-exam-evidence.xlsx"))

with open(os.path.join(OUT, "M14-control-gauge-test.csv"), "w", newline="") as f:
    writer = csv.writer(f, lineterminator="\n")
    writer.writerow([
        "pass_id", "test_truck_weight_kip", "gauge_id",
        "peak_strain_microstrain", "speed_mph", "lane_position", "quality_flag",
    ])
    for idx, strain in enumerate(STRAINS, 1):
        writer.writerow([f"LT-{idx:02d}", f"{TEST_LOAD_KIP:.1f}", "G4",
                         f"{strain:.1f}", "10", "center", "VALID"])

print(
    "  wrote M14 clean exam evidence: "
    f"{len(STRAINS)} passes, {TEST_LOAD_KIP:.0f}-kip test, "
    f"{PERMIT_LOAD_KIP:.0f}-kip request"
)
