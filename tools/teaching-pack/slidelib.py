"""Slide builders for the ACMEJOB.Ai decks.

Visual system: colour says WHO IS SPEAKING.
  Orange  = FOREMAN (the machine)      Blue = a person (Diane, Wes, the county)
That rule holds on every slide of every deck and is the spine of the course.

Typography: display titles in Barlow Condensed (short strings, generous slack);
all body text in Arial, which is metric-predictable everywhere, so a podium
machine without the brand fonts still lays out correctly.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(ROOT, "kit", "logo")
CHARTS = os.path.join(ROOT, "build", "charts")

W, H = 13.333, 7.5

GIRDER = RGBColor(0x1F, 0x3A, 0x5F)
GIRDER_LT = RGBColor(0x3A, 0x7D, 0xBF)
STICKER = RGBColor(0xF2, 0x6B, 0x1D)
REBAR = RGBColor(0x2B, 0x2F, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CONCRETE = RGBColor(0xE6, 0xE8, 0xEA)
RIVER = RGBColor(0x6F, 0x8F, 0xAF)
GREY = RGBColor(0x6B, 0x72, 0x80)
CRIT = RGBColor(0xC0, 0x26, 0x3C)
PAPER = RGBColor(0xF4, 0xF6, 0xF8)
CREAM = RGBColor(0xFF, 0xF6, 0xE5)
PEACH = RGBColor(0xFF, 0xF1, 0xE8)
GOLD = RGBColor(0xF2, 0xC2, 0x30)

DISPLAY = "Barlow Condensed"
BODY = "Arial"


def deck():
    p = Presentation()
    p.slide_width = Inches(W)
    p.slide_height = Inches(H)
    return p


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, prs, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)
    return s


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, lw=1.25):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        s.adjustments[0] = radius
    if fill is not None:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line is not None:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, runs, size=18, color=REBAR, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.0, space_after=6,
         italic=False, margin=0):
    """runs: a string, or a list of (text, {opts}) tuples, or a list of paragraphs
    where each paragraph is a string or a list of run tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tb.text_frame.margin_left = Inches(margin)
    tb.text_frame.margin_right = Inches(margin)
    tb.text_frame.margin_top = Inches(margin)
    tb.text_frame.margin_bottom = Inches(margin)

    paras = runs if isinstance(runs, list) else [runs]
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line
        p.space_after = Pt(space_after)
        items = para if isinstance(para, list) else [(para, {})]
        for t, o in items:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.name = o.get("font", font)
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def bullet_list(slide, x, y, w, h, items, size=17, color=REBAR, gap=13, dot=STICKER):
    """Hanging-indent bullets drawn with real dots, so nothing depends on list styling."""
    cy = y
    for it in items:
        lead, rest = (it if isinstance(it, tuple) else (None, it))
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(cy + 0.105),
                                   Inches(0.085), Inches(0.085))
        d.fill.solid(); d.fill.fore_color.rgb = dot
        d.line.fill.background(); d.shadow.inherit = False
        runs = []
        if lead:
            runs.append((lead, {"bold": True, "color": color}))
        runs.append((rest, {"color": color}))
        est = 0.30 + 0.245 * (len(lead or "") + len(rest)) // int(w * 13.5)
        tb = text(slide, x + 0.26, cy, w - 0.26, est, [runs], size=size,
                  color=color, line=1.12, space_after=0)
        tb.text_frame.word_wrap = True
        n_lines = max(1, -(-(len(lead or "") + len(rest)) // int(w * 6.2)))
        cy += 0.30 * n_lines + gap / 72
    return cy


def logo_mark(slide, x, y, h=0.34, dark=False):
    f = "acmejob-logo-horizontal.png" if not dark else "acmejob-logo-horizontal-dark.png"
    p = os.path.join(LOGO, f)
    if not os.path.exists(p):
        p = os.path.join(LOGO, "acmejob-logo-horizontal.png")
    return slide.shapes.add_picture(p, Inches(x), Inches(y), height=Inches(h))


def picture(slide, name, x, y, w=None, h=None):
    p = name if os.path.isabs(name) else os.path.join(CHARTS, name)
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    return slide.shapes.add_picture(p, Inches(x), Inches(y), **kw)


def fit(name, box_w, box_h):
    """Scale an image to fit a box, returning (x_off_w, x_off_h) sizes in inches."""
    from PIL import Image
    p = name if os.path.isabs(name) else os.path.join(CHARTS, name)
    iw, ih = Image.open(p).size
    ar = iw / ih
    w, h = box_w, box_w / ar
    if h > box_h:
        h, w = box_h, box_h * ar
    return w, h


def place(slide, name, cx, cy, box_w, box_h):
    """Centre an image inside a box."""
    w, h = fit(name, box_w, box_h)
    return picture(slide, name, cx - w / 2, cy - h / 2, w=w)


# ------------------------------------------------------------- layouts ------
def title_slide(prs, meeting, title, subtitle, date):
    s = blank(prs); bg(s, prs, REBAR)
    picture(s, os.path.join(LOGO, "acmejob-mark-dark.png"), 0.95, 1.05, h=1.5)
    text(s, 0.95, 2.85, 8.6, 0.42, meeting.upper(), size=17, color=STICKER,
         font=DISPLAY, bold=True)
    text(s, 0.95, 3.25, 11.4, 1.6, title, size=58, color=WHITE, font=DISPLAY, bold=True,
         line=0.96)
    text(s, 0.95, 5.0, 10.2, 1.0, subtitle, size=22, color=RIVER, line=1.2)
    text(s, 0.95, 6.5, 10.2, 0.4, date, size=14, color=GREY, font=DISPLAY)
    return s


def section(prs, num, title, kicker=None):
    s = blank(prs); bg(s, prs, GIRDER)
    text(s, 1.0, 2.5, 2.0, 1.0, num, size=90, color=STICKER, font=DISPLAY, bold=True, line=0.9)
    text(s, 1.0, 3.6, 11.0, 1.4, title, size=46, color=WHITE, font=DISPLAY, bold=True, line=1.0)
    if kicker:
        text(s, 1.0, 5.0, 10.0, 0.8, kicker, size=19, color=RIVER, line=1.2)
    return s


def statement(prs, line, attrib=None, dark=True, size=44, accent=None):
    s = blank(prs); bg(s, prs, REBAR if dark else WHITE)
    col = WHITE if dark else GIRDER
    text(s, 1.1, 2.3, 11.1, 2.8, line, size=size, color=accent or col, font=DISPLAY,
         bold=True, line=1.06, anchor=MSO_ANCHOR.MIDDLE)
    if attrib:
        text(s, 1.1, 5.25, 11.1, 0.5, attrib, size=17,
             color=RIVER if dark else GREY, italic=True)
    return s


def head(slide, title, kicker=None, color=GIRDER):
    text(slide, 0.72, 0.5, 11.9, 0.85, title, size=38, color=color, font=DISPLAY,
         bold=True, line=0.98)
    y = 1.32
    if kicker:
        text(slide, 0.72, y, 11.9, 0.5, kicker, size=17, color=GREY, line=1.15)
        y += 0.62
    return y


def content(prs, title, kicker=None):
    s = blank(prs); bg(s, prs, WHITE)
    y = head(s, title, kicker)
    return s, y


WARNINGS = []


def text_height(paras, box_w, size, line=1.18, space_after=7, font=BODY):
    """Estimate rendered height in inches. Calibrated against LibreOffice renders of
    Arial: usable characters per inch of box width is about 144 / point-size."""
    cpi = 144.0 / size
    if font == DISPLAY:
        cpi *= 1.22
    per_line = max(8, int(box_w * cpi))
    total_lines, n = 0, 0
    for p in (paras if isinstance(paras, list) else [paras]):
        s = p if isinstance(p, str) else "".join(t for t, _ in p)
        total_lines += max(1, -(-len(s) // per_line))
        n += 1
    return total_lines * size * line / 72.0 + (n - 1) * space_after / 72.0


def card(slide, x, y, w, h, label, body, tint=PAPER, edge=None, label_color=None,
         body_size=16, label_size=14, mono=False, name=None):
    paras = body if isinstance(body, list) else [body]
    need = 0.72 + text_height(paras, w - 0.6, body_size) + 0.34
    if need > h + 0.02:
        WARNINGS.append(f"CARD OVERFLOW {name or label[:28]!r}: h={h:.2f} needs {need:.2f}")
    rect(slide, x, y, w, h, fill=tint, line=edge or tint, radius=0.035)
    text(slide, x + 0.3, y + 0.24, w - 0.6, 0.32, label, size=label_size,
         color=label_color or GREY, font=DISPLAY, bold=True)
    text(slide, x + 0.3, y + 0.72, w - 0.6, max(0.3, h - 1.0), paras,
         size=body_size, color=REBAR, line=1.18, space_after=7,
         font="Courier New" if mono else BODY)


def speaker_card(slide, x, y, w, h, who, body, machine=False, body_size=16):
    """The motif: orange for FOREMAN, blue for a person."""
    tint = PEACH if machine else RGBColor(0xEE, 0xF3, 0xF8)
    edge = STICKER if machine else GIRDER_LT
    card(slide, x, y, w, h, who, body, tint=tint, edge=edge,
         label_color=STICKER if machine else GIRDER, body_size=body_size)


def activity(prs, title, minutes, steps, note=None):
    s = blank(prs); bg(s, prs, WHITE)
    rect(s, 0, 0, W, 1.65, fill=GOLD)
    text(s, 0.72, 0.36, 9.4, 0.5, "IN CLASS — DO THIS NOW", size=15, color=REBAR,
         font=DISPLAY, bold=True)
    text(s, 0.72, 0.78, 9.4, 0.7, title, size=34, color=REBAR, font=DISPLAY, bold=True,
         line=1.0)
    rect(s, 10.9, 0.42, 1.72, 0.82, fill=REBAR, radius=0.12)
    text(s, 10.9, 0.52, 1.72, 0.62, f"{minutes} min", size=26, color=GOLD,
         font=DISPLAY, bold=True, align=PP_ALIGN.CENTER)
    y = 2.15
    for i, st in enumerate(steps, 1):
        n = rect(s, 0.75, y, 0.46, 0.46, fill=GIRDER, radius=0.5)
        text(s, 0.75, y + 0.07, 0.46, 0.34, str(i), size=19, color=WHITE,
             font=DISPLAY, bold=True, align=PP_ALIGN.CENTER)
        text(s, 1.42, y + 0.02, 10.9, 0.8, st, size=18, color=REBAR, line=1.14)
        y += 0.52 + 0.30 * max(0, (len(st) - 1) // 92)
    if note:
        text(s, 0.75, 6.45, 11.8, 0.6, note, size=15, color=GREY, italic=True, line=1.15)
    return s


def table_slide(prs, title, rows, widths, kicker=None, size=14, head_fill=GIRDER,
                row_colors=None, y0=None, col_align=None):
    s, y = content(prs, title, kicker)
    y = y0 or y
    tw = sum(widths)
    x = (W - tw) / 2
    rows_n, cols_n = len(rows), len(rows[0])
    rh = 0.44
    gt = s.shapes.add_table(rows_n, cols_n, Inches(x), Inches(y),
                            Inches(tw), Inches(rh * rows_n)).table
    gt.first_row = True
    for ci, cw in enumerate(widths):
        gt.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        gt.rows[ri].height = Inches(rh)
        for ci, val in enumerate(row):
            c = gt.cell(ri, ci)
            c.margin_left = Inches(0.13); c.margin_right = Inches(0.1)
            c.margin_top = Inches(0.05); c.margin_bottom = Inches(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (PP_ALIGN.CENTER if col_align and col_align[ci] == "c"
                           else PP_ALIGN.RIGHT if col_align and col_align[ci] == "r"
                           else PP_ALIGN.LEFT)
            r = p.add_run(); r.text = str(val)
            f = r.font
            f.size = Pt(size if ri else size)
            f.name = DISPLAY if ri == 0 else BODY
            f.bold = ri == 0
            f.color.rgb = WHITE if ri == 0 else REBAR
            c.fill.solid()
            if ri == 0:
                c.fill.fore_color.rgb = head_fill
            elif row_colors and row_colors.get(ri):
                c.fill.fore_color.rgb = row_colors[ri]
                if row_colors[ri] == CRIT:
                    f.color.rgb = WHITE; f.bold = True
            else:
                c.fill.fore_color.rgb = WHITE if ri % 2 else PAPER
    return s, y + rh * rows_n


def image_slide(prs, title, img, kicker=None, caption=None, box_h=4.5):
    s, y = content(prs, title, kicker)
    w, h = fit(img, 11.6, box_h)
    picture(s, img, (W - w) / 2, y + 0.12, w=w)
    if caption:
        text(s, 0.72, y + h + 0.32, 11.9, 0.7, caption, size=16, color=GREY, line=1.18)
    return s


def closer(prs, lines, title="Before you go"):
    s = blank(prs); bg(s, prs, GIRDER)
    text(s, 0.95, 0.75, 11.4, 0.8, title, size=38, color=WHITE, font=DISPLAY, bold=True)
    y = 1.9
    for lead, body in lines:
        text(s, 0.95, y, 3.0, 0.5, lead, size=19, color=STICKER, font=DISPLAY, bold=True)
        text(s, 4.1, y - 0.04, 8.3, 0.9, body, size=18, color=WHITE, line=1.15)
        y += 0.95
    return s


def save(prs, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    print("  wrote", os.path.basename(path))
    if WARNINGS:
        for w in WARNINGS:
            print("   !", w)
        print(f"   ! {len(WARNINGS)} layout warning(s)")
    else:
        print("   no layout warnings")
    WARNINGS.clear()
