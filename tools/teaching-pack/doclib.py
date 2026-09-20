"""Shared builders for branded ACMEJOB.Ai / Kinnick County / FOREMAN documents."""
import os, sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(ROOT, "kit", "logo")
BUILD = os.path.join(ROOT, "build")
FACTS = json.load(open(os.path.join(BUILD, "facts.json")))

GIRDER = RGBColor(0x1F, 0x3A, 0x5F)
STICKER = RGBColor(0xF2, 0x6B, 0x1D)
REBAR = RGBColor(0x2B, 0x2F, 0x36)
RIVER = RGBColor(0x6F, 0x8F, 0xAF)
RUST = RGBColor(0x9C, 0x4A, 0x2F)
GREY = RGBColor(0x6B, 0x72, 0x80)

HEAD_FONT = "Barlow Condensed"
BODY_FONT = "Barlow"
MONO_FONT = "IBM Plex Mono"


# --------------------------------------------------------------- helpers ----
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def cell_margins(table, top=60, bottom=60, left=110, right=110):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        b.append(e)
    tblPr.append(b)


def bottom_rule(par, color="1F3A5F", size=8):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), "4")
    e.set(qn("w:color"), color)
    bdr.append(e)
    pPr.append(bdr)


def box(par, fill=None, color="D5D9DE", size=6, space=8):
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for tag in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), color)
        bdr.append(e)
    pPr.append(bdr)
    if fill:
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), fill)
        pPr.append(sh)


def page_field(par):
    for txt, fld in (("", "PAGE"), (" of ", None), ("", "NUMPAGES")):
        if fld is None:
            r = par.add_run(txt)
            r.font.size = Pt(8)
            r.font.color.rgb = GREY
            r.font.name = BODY_FONT
            continue
        r = par.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        r.font.name = BODY_FONT
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {fld} "
        fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "end")
        r._r.append(fc); r._r.append(it); r._r.append(fc2)


def run(par, text, size=10, bold=False, italic=False, color=None, font=BODY_FONT,
        caps=False, space=None):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), font)
    if color is not None:
        r.font.color.rgb = color
    if caps:
        c = OxmlElement("w:caps"); c.set(qn("w:val"), "1"); rPr.append(c)
    if space:
        s = OxmlElement("w:spacing"); s.set(qn("w:val"), str(int(space * 20))); rPr.append(s)
    return r


def para(doc, text="", size=10, bold=False, italic=False, color=None, font=BODY_FONT,
         align=None, before=0, after=6, style=None, line=1.15, caps=False, space=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, italic, color, font, caps, space)
    return p


def h1(doc, text, color=None):
    return para(doc, text, size=21, bold=True, color=color or GIRDER, font=HEAD_FONT,
                before=2, after=6, line=1.0)


def h2(doc, text, color=None, before=13):
    p = para(doc, text, size=12.5, bold=True, color=color or GIRDER, font=HEAD_FONT,
             before=before, after=4, line=1.0, caps=True, space=0.6)
    return p


def bullets(doc, items, size=10, indent=0.24, after=3):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.left_indent = Inches(indent)
        if isinstance(it, tuple):
            run(p, it[0], size, bold=True)
            run(p, it[1], size)
        else:
            run(p, it, size)


def numbers(doc, items, size=10, indent=0.24):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.left_indent = Inches(indent)
        if isinstance(it, tuple):
            run(p, it[0], size, bold=True)
            run(p, it[1], size)
        else:
            run(p, it, size)


def table(doc, rows, widths, header=True, size=9.5, head_fill="1F3A5F",
          head_ink=RGBColor(0xFF, 0xFF, 0xFF), zebra="F4F6F8", align=None, font=BODY_FONT):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell_margins(t)
    total = Inches(sum(widths))
    t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.width = Inches(widths[ci])
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if align and align[ci] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align and align[ci] == "c":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            is_head = header and ri == 0
            run(p, str(val), size, bold=is_head,
                color=head_ink if is_head else None,
                font=HEAD_FONT if is_head else font)
            if is_head:
                shade(c, head_fill)
            elif zebra and ri % 2 == 0:
                shade(c, zebra)
    return t


def callout(doc, label, text, fill="FFF6E5", edge="F2C230", size=9.5, after=8):
    p = para(doc, "", after=after, before=6)
    box(p, fill=fill, color=edge, size=8, space=6)
    if label:
        run(p, label + "  ", size, bold=True, color=REBAR, font=HEAD_FONT, caps=True, space=0.5)
    run(p, text, size)
    return p


def handwriting(doc, text, size=13, after=6, before=4, indent=0.2):
    p = para(doc, "", before=before, after=after)
    p.paragraph_format.left_indent = Inches(indent)
    run(p, text, size, font="Caveat", color=GIRDER)
    return p


def mono_block(doc, lines, size=8.5, fill="F4F6F8", edge="D5D9DE"):
    """Monospace block. Each line gets the fill; only the outer edges get a border."""
    n = len(lines)
    for i, ln in enumerate(lines):
        p = para(doc, "", before=0, after=0, line=1.15)
        p.paragraph_format.left_indent = Inches(0.06)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        sides = [("left", True), ("right", True),
                 ("top", i == 0), ("bottom", i == n - 1)]
        for tag, on in sides:
            e = OxmlElement(f"w:{tag}")
            e.set(qn("w:val"), "single" if on else "nil")
            e.set(qn("w:sz"), "4" if on else "0")
            e.set(qn("w:space"), "4")
            e.set(qn("w:color"), edge)
            bdr.append(e)
        pPr.append(bdr)
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill)
        pPr.append(sh)
        run(p, ln if ln else " ", size, font=MONO_FONT, color=REBAR)
    return p


# ------------------------------------------------------------ letterheads ---
def _base(margins=(0.85, 0.8, 0.8, 0.8)):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    for s in d.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = Inches(margins[0]); s.bottom_margin = Inches(margins[1])
        s.left_margin = Inches(margins[2]); s.right_margin = Inches(margins[3])
        s.header_distance = Inches(0.4); s.footer_distance = Inches(0.35)
    return d


def _footer(doc, left_text, ink=GREY):
    f = doc.sections[0].footer
    p = f.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    # tab stop at right margin
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "right"); tab.set(qn("w:pos"), str(int(6.9 * 1440)))
    tabs.append(tab); pPr.append(tabs)
    run(p, left_text, 8, color=ink)
    p.add_run("\t")
    page_field(p)
    return p


def acme_doc(subtitle=None, footer="Otter Bend Lift Bridge  ·  Kinnick County Project 27-114"):
    """ACMEJOB.Ai letterhead."""
    d = _base()
    hdr = d.sections[0].header
    p = hdr.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(LOGO, "acmejob-logo-horizontal.png"), height=Inches(0.34))
    if subtitle:
        p2 = hdr.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(3)
        run(p2, subtitle, 8.5, color=GREY, font=HEAD_FONT, caps=True, space=1.0)
        bottom_rule(p2, "C9CFD6", 6)
    else:
        bottom_rule(p, "C9CFD6", 6)
    _footer(d, footer)
    return d


def foreman_doc(title, generated="2027-02-02 08:14 CST", job="27-114 Otter Bend"):
    """FOREMAN output. Deliberately over-confident machine formatting."""
    d = _base(margins=(0.75, 0.8, 0.8, 0.8))
    hdr = d.sections[0].header
    p = hdr.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(os.path.join(LOGO, "foreman-mark-wide.png"), height=Inches(0.3))
    p2 = hdr.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(3)
    run(p2, f"GENERATED  {generated}   ·   JOB {job}   ·   MODEL fm-4.2-turbo",
        7.5, color=GREY, font=MONO_FONT)
    bottom_rule(p2, "F26B1D", 10)
    _footer(d, "FOREMAN automated output  ·  ACMEJOB.Ai internal  ·  Not reviewed by a licensed engineer")
    h1(d, title, color=REBAR)
    return d


def county_doc(subtitle="Department of Public Works"):
    d = _base()
    hdr = d.sections[0].header
    t = hdr.add_table(rows=1, cols=2, width=Inches(6.9))
    no_borders(t)
    cell_margins(t, 0, 0, 0, 0)
    t.cell(0, 0).width = Inches(0.75)
    t.cell(0, 1).width = Inches(6.15)
    t.cell(0, 0).paragraphs[0].add_run().add_picture(
        os.path.join(LOGO, "kinnick-county-seal.png"), height=Inches(0.62))
    c = t.cell(0, 1).paragraphs[0]
    c.paragraph_format.space_before = Pt(6)
    run(c, "KINNICK COUNTY, MINNESOTA", 14, bold=True, color=RUST, font=HEAD_FONT, space=1.2)
    c2 = t.cell(0, 1).add_paragraph()
    run(c2, subtitle, 9, color=GREY, font=HEAD_FONT, caps=True, space=0.8)
    hp = hdr.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    endp = hdr.add_paragraph()
    endp.paragraph_format.space_before = Pt(2)
    endp.paragraph_format.space_after = Pt(2)
    bottom_rule(endp, "9C4A2F", 8)
    _footer(d, "Kinnick County Public Works  ·  114 Court St, Otter Bend, MN")
    return d


def course_doc(meeting, title, kind="HANDOUT"):
    """Plain course handout: no firm branding, for things that are course scaffolding."""
    d = _base()
    hdr = d.sections[0].header
    p = hdr.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run(p, f"SEIS 201  ·  AI FOR CE/ME ENGINEERS", 8.5, color=GREY, font=HEAD_FONT, caps=True, space=1.0)
    p.add_run("\t\t")
    run(p, f"{meeting}  ·  {kind}", 8.5, color=GREY, font=HEAD_FONT, caps=True, space=1.0)
    bottom_rule(p, "C9CFD6", 6)
    _footer(d, f"SEIS 201  ·  {meeting}")
    h1(d, title)
    return d


def memo_block(doc, to, frm, date, re_, cc=None):
    rows = [("TO", to), ("FROM", frm), ("DATE", date)]
    if cc:
        rows.append(("CC", cc))
    rows.append(("RE", re_))
    t = doc.add_table(rows=len(rows), cols=2)
    no_borders(t)
    cell_margins(t, 10, 10, 0, 60)
    t.autofit = False
    for i, (k, v) in enumerate(rows):
        t.cell(i, 0).width = Inches(0.62)
        t.cell(i, 1).width = Inches(6.28)
        p = t.cell(i, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run(p, k, 8.5, bold=True, color=GREY, font=HEAD_FONT, caps=True, space=0.8)
        p2 = t.cell(i, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        run(p2, v, 10, bold=(k == "RE"))
    p = para(doc, "", after=9, before=1, line=1.0)
    bottom_rule(p, "C9CFD6", 6)
    return t


def save(doc, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    print("  wrote", os.path.relpath(path, BUILD))
    return path
