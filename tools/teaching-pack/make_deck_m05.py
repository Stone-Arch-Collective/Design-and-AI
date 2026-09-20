"""M05 spoiler-safe student deck and gated instructor reveal."""
import os
import sys
from io import BytesIO

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from slidelib import *


def add_notes(prs, notes):
    if len(prs.slides) != len(notes):
        raise ValueError(f"notes ({len(notes)}) != slides ({len(prs.slides)})")
    for slide, note in zip(prs.slides, notes):
        slide.notes_slide.notes_text_frame.text = note.strip()


def select_slides(source, numbers):
    buf = BytesIO()
    source.save(buf)
    buf.seek(0)
    result = Presentation(buf)
    keep = set(numbers)
    for i in reversed(range(len(result.slides))):
        if i + 1 not in keep:
            slide_id = result.slides._sldIdLst[i]
            result.part.drop_rel(slide_id.rId)
            del result.slides._sldIdLst[i]
    return result


prs = deck()

# 1
title_slide(prs, "Meeting 05  ·  Unit 1: Hired",
            "The 3 a.m. alarm.",
            "Trace an agent from live stream to action—one step at a time",
            "Tuesday, February 16, 2027  ·  SEIS 201")

# 2
s, y = content(prs, "At 3:07 a.m., FOREMAN drafted this for the county",
               "The send was held for Diane's approval")
speaker_card(s, .72, y + .15, 11.87, 2.2, "FOREMAN v4.2 — EMERGENCY",
             ["“Threshold exceeded. Recommend immediate closure of the Otter Bend Lift Bridge "
              "pending inspection.”"], machine=True, body_size=23)
speaker_card(s, .72, y + 2.7, 11.87, 1.55, "DIANE HALVORSEN, PE",
             ["“The draft did not leave this office. Show me how it got from the feed to that sentence.”"],
             body_size=18)

# 3
s, y = content(prs, "Your job is not to guess what crossed the bridge",
               "Your job is to reconstruct what the agent did")
card(s, .72, y + .2, 11.87, 1.7, "THE HUNT",
     ["Was the alarm a recorded exceedance, a pipeline artifact, a threshold mistake, an "
      "averaging mistake—or a correct warning followed by an unsupported action?"],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=18)
bullet_list(s, .95, y + 2.15, 11.4, 2.4, [
    ("Artifacts  ", "FOREMAN-overnight-run.log, threshold_config.csv, feed.csv."),
    ("Eight hours  ", "choose which handoffs to verify; hours are not points."),
    ("Boundary  ", "a recorded number is not yet an explanation of a physical event."),
], size=18, dot=GIRDER_LT)

# 4
section(prs, "1", "What makes it an agent?",
        "Text generation plus tools, state, a plan, and permission to act")

# 5
s, y = content(prs, "An AI agent runs a loop",
               "It observes a state, chooses a next step, and changes the state")
rows = [
    ["Loop part", "Question", "M05 example"],
    ["Plan", "What should happen next?", "Poll → normalize → summarize → test → draft"],
    ["Tool call", "What operation leaves the model?", "Query feed; apply offset; compute mean"],
    ["Observation", "What came back?", "Eight rows; max 286.1 µε; threshold true"],
    ["Decision", "What does it mean / what next?", "Choose severity and next action"],
]
gt = s.shapes.add_table(len(rows), 3, Inches(.72), Inches(y + .2),
                        Inches(11.87), Inches(3.5)).table
for ci, width in enumerate((2.0, 3.2, 6.67)):
    gt.columns[ci].width = Inches(width)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        c = gt.cell(ri, ci)
        c.margin_left = Inches(.12); c.margin_right = Inches(.08)
        c.fill.solid(); c.fill.fore_color.rgb = GIRDER if ri == 0 else (WHITE if ri % 2 else PAPER)
        p = c.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.name = DISPLAY if ri == 0 else BODY
        r.font.size = Pt(15 if ri == 0 else 14)
        r.font.bold = ri == 0
        r.font.color.rgb = WHITE if ri == 0 else REBAR
card(s, .72, y + 3.95, 11.87, 1.32, "IMPORTANT",
     ["The loop label tells you what kind of step it is. It does not tell you the step is justified."],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=14)

# 6
s, y = content(prs, "Tools make outputs consequential",
               "A model writes tokens. An agent can read files, calculate, draft, queue, or send.")
speaker_card(s, .72, y + .2, 5.85, 3.3, "MODEL",
             ["Input: text", "", "Output: text", "", "No direct change outside the conversation."],
             body_size=19)
card(s, 6.72, y + .2, 5.87, 3.3, "AGENT",
     ["Input: state + goal", "", "Output: tool calls + changed state", "",
      "Its permissions determine what the output can do."],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=19)
text(s, .72, y + 3.9, 11.9, .7,
     "The engineering question moves from “Is the sentence right?” to “Is every handoff justified?”",
     size=20, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 7
s, y = content(prs, "A sensor stream is ordered evidence",
               "New rows arrive; time and identity are part of every value")
bullet_list(s, .95, y + .25, 11.4, 3.35, [
    ("Timestamp  ", "when the reading belongs—not when a report was drafted."),
    ("Sensor identity  ", "where and how the value was measured."),
    ("Raw field  ", "what the instrument reported."),
    ("Transformation  ", "baseline correction, filtering, aggregation, or unit conversion."),
    ("Quality metadata  ", "what the pipeline knows about missing, suspect, or repeated data."),
], size=18, dot=GIRDER_LT)
card(s, .72, y + 3.92, 11.87, 1.35, "AUDIT RULE",
     ["Never let a transformed value travel without the field name, time window, and sample count."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=15)

# 8
s, y = content(prs, "The threshold is a rule inside the agent",
               "Learned system outside; explicit comparison inside")
rect(s, .72, y + .25, 11.87, 1.25, fill=PEACH, line=STICKER, radius=.05)
text(s, .72, y + .5, 11.87, .7,
     "IF corrected 10-minute mean > 250 µε  →  WARN-250 = TRUE",
     size=29, color=STICKER, font="Courier New", bold=True, align=PP_ALIGN.CENTER)
bullet_list(s, .95, y + 1.95, 11.4, 2.3, [
    ("Rule-based  ", "the same input gives the same comparison result."),
    ("Not self-explanatory  ", "the threshold source, field, window, and action still need checking."),
    ("A true condition is not a complete disposition. ", "The action rule is a separate handoff."),
], size=18, dot=STICKER)

# 9
s, y = content(prs, "Autonomy changes when review can happen",
               "FOREMAN had one overnight hour and permission to prepare—not send—a draft")
card(s, .72, y + .2, 5.85, 3.35, "OVERNIGHT AUTONOMY",
     ["Poll live data", "Call approved analysis tools", "Open an internal flag",
      "Prepare a draft", "Return unused project hours"],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=17)
speaker_card(s, 6.72, y + .2, 5.87, 3.35, "HUMAN CHECKPOINT",
             ["Review evidence", "Choose external action", "Approve or reject county communication",
              "Remain responsible for the disposition"], body_size=17)
text(s, .72, y + 3.95, 11.9, .7,
     "A queued draft is safer than an auto-send—and can still contain an unsupported recommendation.",
     size=18, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 10
activity(prs, "Label the overnight run", 13, [
    "Open FOREMAN-overnight-run.log and H5-03.",
    "Mark each substantive line PLAN, TOOL, OBSERVATION, or DECISION.",
    "Draw an arrow from every observation to the decision that uses it.",
    "Circle the first decision whose evidence you cannot yet reproduce.",
], note="Do not decide whether the physical event was real. Trace the action loop first.")

# 11
activity(prs, "Audit the pipeline, one boundary at a time", 22, [
    "Filter feed.csv to 2027-02-16 03:00. Follow raw strain through baseline correction.",
    "Recompute the 10-minute mean and record the number of samples in the window.",
    "Evaluate WARN-250 from threshold_config.csv.",
    "Compare the action FOREMAN took with WARN-250 and CLOSE-DRAFT. Mark the first mismatch.",
], note="Use evidence for every claim: file, row, field, rule, and arithmetic.")

# 12
s, y = content(prs, "Keep three questions separate",
               "They require different evidence")
card(s, .72, y + .2, 3.75, 3.45, "1  RECORDED?",
     ["Does the value exist in feed.csv after the stated transformation?"],
     tint=PAPER, edge=CONCRETE, label_color=GIRDER, body_size=18)
card(s, 4.79, y + .2, 3.75, 3.45, "2  TRIGGERED?",
     ["Does the configured threshold evaluate true on that value and window?"],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=18)
card(s, 8.86, y + .2, 3.75, 3.45, "3  PHYSICAL?",
     ["What caused the reading, and what action does that cause justify?"],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=18)
text(s, .72, y + 4.0, 11.9, .7,
     "One file may answer the first two and still be insufficient for the third.",
     size=20, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 13
statement(prs, "Where did “look at this”\nbecome “close it”?",
          "Write your three-sentence note before the reveal.", dark=False)

# 14
s, y = content(prs, "The recorded warning is reproducible",
               "The available evidence supports an internal flag")
rows = [
    ["Check", "Result"],
    ["Corrected G6 at 03:00", "286.1 µε exists in feed.csv"],
    ["Baseline correction", "187 µε removed; did not create the exceedance"],
    ["10-minute mean", "1 sample; mean equals that single reading"],
    ["WARN-250", "286.1 > 250.0 → TRUE"],
]
gt = s.shapes.add_table(len(rows), 2, Inches(.72), Inches(y + .25),
                        Inches(11.87), Inches(3.0)).table
gt.columns[0].width = Inches(4.0); gt.columns[1].width = Inches(7.87)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        c = gt.cell(ri, ci)
        c.fill.solid(); c.fill.fore_color.rgb = GIRDER if ri == 0 else (WHITE if ri % 2 else PAPER)
        p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = val
        r.font.name = DISPLAY if ri == 0 else BODY; r.font.size = Pt(17)
        r.font.bold = ri == 0; r.font.color.rgb = WHITE if ri == 0 else REBAR
card(s, .72, y + 3.65, 11.87, 1.35, "WHAT THAT DOES NOT YET ESTABLISH",
     ["The physical cause of the reading, persistence at the next poll, or an immediate-closure disposition."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=16)

# 15
s, y = content(prs, "The unsupported jump is in the action",
               "FOREMAN evaluated one rule and acted as if it had satisfied another")
speaker_card(s, .72, y + .2, 5.85, 3.55, "WARN-250 ACTUALLY SAYS",
             ["Open internal review flag.", "Repeat at next poll.", "Inspect neighbours.",
              "Require human review before external communication."], body_size=17)
card(s, 6.72, y + .2, 5.87, 3.55, "FOREMAN DID",
     ["Declared immediate closure safest.", "Created emergency county draft.",
      "Said no further checks were required."],
     tint=PEACH, edge=STICKER, label_color=STICKER, body_size=17)
text(s, .72, y + 4.1, 11.9, .6,
     "Correct comparison. Unsupported disposition.",
     size=22, color=STICKER, bold=True, align=PP_ALIGN.CENTER)

# 16
s, y = content(prs, "A one-point average is still one point",
               "The label “10-minute mean” sounds more robust than the operation was")
rect(s, .72, y + .3, 11.87, 1.35, fill=PAPER, line=CONCRETE, radius=.05)
text(s, .72, y + .62, 11.87, .7, "mean([286.1]) = 286.1", size=34,
     color=GIRDER, font="Courier New", bold=True, align=PP_ALIGN.CENTER)
bullet_list(s, .95, y + 2.05, 11.4, 2.25, [
    ("Averaging did not manufacture the value. ", "It also did not add repeat evidence."),
    ("The sample count is decision-relevant metadata. ", "Window length alone does not tell you how much data entered."),
    ("The next poll matters. ", "Persistence was required and unavailable at 03:07."),
], size=18, dot=GIRDER_LT)

# 17
s, y = content(prs, "The smallest defensible action",
               "Match the action to the evidence you actually have")
card(s, .72, y + .2, 11.87, 3.35, "MODEL NOTE — AFTER STUDENTS WRITE",
     ["The 03:00 G6 value in the corrected feed exceeds the 250 µε warning threshold, and I can "
      "reproduce that comparison.",
      "The available window contains one G6 reading, so it does not support FOREMAN's immediate-"
      "closure recommendation or establish what caused the exceedance.",
      "Keep the internal flag open, obtain the next poll, compare adjacent gauges, and require "
      "licensed-engineer review before any county message."],
     tint=RGBColor(0xEE, 0xF3, 0xF8), edge=GIRDER_LT, label_color=GIRDER, body_size=17)
text(s, .72, y + 3.9, 11.9, .75,
     "This is not “do nothing.” It is a proportionate checkpoint.",
     size=19, color=GIRDER, bold=True, align=PP_ALIGN.CENTER)

# 18
s, y = content(prs, "The question M05 cannot answer",
               "Was the reading caused by a real load or by an artifact?")
bullet_list(s, .95, y + .35, 11.4, 3.0, [
    ("You have checked  ", "the row, correction, averaging, and threshold comparison."),
    ("You have not checked  ", "the event against an independent source."),
    ("Do not erase it  ", "because it looks unusual."),
    ("Do not close the bridge  ", "because one threshold comparison returned true."),
], size=19, dot=STICKER)
card(s, .72, y + 3.78, 11.87, 1.4, "THURSDAY",
     ["Audit FOREMAN's reversal using the distribution and the traffic record."],
     tint=CREAM, edge=GOLD, label_color=REBAR, body_size=16)

# 19
s = blank(prs); bg(s, prs, WHITE); head(s, "Mask off — two minutes")
bullet_list(s, .95, 1.9, 11.4, 3.5, [
    ("An agent is a chain of handoffs. ", "A correct early step does not validate the later action."),
    ("Rules can sit inside learned systems. ", "The 250 µε comparison is the M01 kind of rule."),
    ("Autonomy spends review opportunities. ", "Put checkpoints before external or irreversible action."),
    ("Diane is sharper because the firm's name was almost attached. ", "Her concern is not the answer key."),
], size=18, dot=GIRDER_LT)

# 20
closer(prs, [
    ("Hand in", "H5-04 note v1: supports, does not support, next checkpoint"),
    ("HW3", "Audit row counts and duplicated timestamps; compute peak corrected strain by timestamp"),
    ("Thursday", "FOREMAN reverses itself; audit the distribution, then hand the gauge file to Wes"),
], title="Before Thursday")

student = select_slides(prs, range(1, 14))
add_notes(student, [
    "[SAY] “At 3:07 this morning, FOREMAN stopped being a chat window and became part of an operational workflow.”",
    "[SAY] “The draft did not go out. Do not reward the checkpoint by assuming the draft was right or wrong. Trace it.”",
    "[SAY] “Your eight hours are a prioritization constraint, not points. Every claim today needs a file, row, field, rule, or calculation.”",
    "[SAY] “Agentic means the system can choose and execute next steps through tools. That expands both usefulness and failure surface.”",
    "[SAY] “Mark the kind of step before judging it. A decision can be formatted cleanly and still outrun its observation.”",
    "[SAY] “Tools are where generated text becomes changed state. Permissions and checkpoints are part of the engineering design.”",
    "[SAY] “Time, sensor identity, field name, and quality metadata travel with the value. A number alone is not a stream record.”",
    "[SAY] “This comparison is deterministic. Finding TRUE answers only one question: did this input cross this boundary?”",
    "[SAY] “Overnight autonomy moved review later. It did not remove responsibility from the human who approves the action.”",
    "[ACTIVITY] Keep the log on screen. Confirm labels only; do not confirm the first unsupported decision.",
    "[ACTIVITY] Require evidence at each boundary. Do not say whether 286.1 is real, artifact, or physical event.",
    "[SAY] “Recorded, triggered, and physical are three claims. Keep them separate in your note.”",
    "[SAY] “Write before I open the reveal. Name the exact handoff where the evidence stopped supporting the action.”",
])
save(student, os.path.join(ROOT, "build", "M05", "M05-student.pptx"))

reveal = select_slides(prs, range(14, 21))
add_notes(reveal, [
    "Open after students commit to a first unsupported line. Reproduce the four checks in order.",
    "The designed error is the action jump, not arithmetic. Read WARN-250 and CLOSE-DRAFT aloud.",
    "Ask how many values went into the mean before showing one. Window duration is not sample count.",
    "Accept equivalent notes that separate recorded exceedance from cause and choose a proportionate checkpoint.",
    "Hold the M06 boundary. Do not disclose the traffic match or call the event a sensor fault.",
    "Mask off the system-design lessons. Diane's concern should not function as the answer key.",
    "Assign the duplicate audit without naming the block. Preview the Wes handoff, not its outcome.",
])
save(reveal, os.path.join(ROOT, "build", "M05", "M05-instructor-reveal.pptx"))
