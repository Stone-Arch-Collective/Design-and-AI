"""Deterministic M13 late handoff and silently-ingested pipeline evidence."""
import csv
import os
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "build", "M13")
os.makedirs(OUT, exist_ok=True)

start = datetime(2027, 3, 14, 18, 0)
offset_tokens = [
    "0.18 ft", "0.21 ft", "130 mm", "0.20 ft", "0.19 ft", "0.23 ft",
    "95 mm", "0.22 ft", "0.18 ft", "0.24 ft", "0.20 ft", "0.21 ft",
]
raw_strain = [118.0, 121.0, 123.0, 126.0, 129.0, 132.0, 135.0, 137.0, 140.0, 143.0, 145.0, 148.0]
temperatures = [29.0, 30.0, 31.0, 32.0, None, 34.0, 35.0, 36.0, 37.0, None, 39.0, 40.0]

sensor_rows = []
weather_rows = []
for i, (offset, strain, temp) in enumerate(zip(offset_tokens, raw_strain, temperatures), 1):
    timestamp = (start + timedelta(hours=i - 1)).strftime("%Y-%m-%d %H:%M CST")
    record_id = f"OB-{100 + i}"
    sensor_rows.append([record_id, timestamp, "G4", f"{strain:.1f}", offset, "OK"])
    if temp is not None:
        weather_rows.append([timestamp, f"{temp:.1f}", "Kinnick County Airport", "FINAL"])


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


wb = Workbook()
sensor = wb.active
sensor.title = "SENSOR_READINGS"
sensor.append([
    "record_id", "timestamp_cst", "sensor_id", "raw_strain_microstrain",
    "lateral_offset", "quality_flag",
])
for row in sensor_rows:
    sensor.append(row)
style_sheet(sensor, [14, 24, 12, 26, 18, 14])

weather = wb.create_sheet("WEATHER_STATION")
weather.append(["timestamp_cst", "air_temperature_f", "station", "record_status"])
for row in weather_rows:
    weather.append(row)
style_sheet(weather, [24, 22, 28, 18])

wb.properties.title = "Otter Bend M13 Wes handoff"
wb.properties.subject = "Two delivered sheets from a four-sheet handoff"
wb.properties.creator = "Wes Tanaka, EIT"
wb.save(os.path.join(OUT, "wes_handoff.xlsx"))

# Reproduce FOREMAN's silent path. Numeric tokens are accepted under a default-foot schema;
# an inner join emits only rows with exact weather matches.
weather_by_time = {row[0]: float(row[1]) for row in weather_rows}
pipeline_rows = []
for record_id, timestamp, sensor_id, strain_text, offset_token, quality in sensor_rows:
    if timestamp not in weather_by_time:
        continue
    strain = float(strain_text)
    temperature = weather_by_time[timestamp]
    parsed_offset = float(offset_token.split()[0])
    thermal_adjustment = 0.42 * (temperature - 32.0)
    alignment_adjustment = 8.0 * parsed_offset
    model_input = strain - thermal_adjustment + alignment_adjustment
    pipeline_rows.append({
        "record_id": record_id,
        "timestamp_cst": timestamp,
        "raw_strain_microstrain": f"{strain:.2f}",
        "air_temperature_f": f"{temperature:.1f}",
        "source_lateral_offset": offset_token,
        "parsed_offset": f"{parsed_offset:.3f}",
        "schema_unit": "ft",
        "thermal_adjustment_microstrain": f"{thermal_adjustment:.2f}",
        "alignment_adjustment_microstrain": f"{alignment_adjustment:.2f}",
        "model_input_microstrain": f"{model_input:.2f}",
    })

with open(os.path.join(OUT, "M13-pipeline-output.csv"), "w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(pipeline_rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(pipeline_rows)

log_rows = [
    ["step", "time_cst", "operation", "input_records", "output_records", "status", "message"],
    ["01", "2027-03-16 20:01:04", "open SENSOR_READINGS", "—", "12", "COMPLETE", "schema accepted"],
    ["02", "2027-03-16 20:01:05", "parse lateral_offset numeric token", "12", "12", "COMPLETE", "default schema unit=ft"],
    ["03", "2027-03-16 20:01:06", "open WEATHER_STATION", "—", "10", "COMPLETE", "schema accepted"],
    ["04", "2027-03-16 20:01:07", "exact timestamp inner join", "12 + 10", "10", "COMPLETE", "joined dataset emitted"],
    ["05", "2027-03-16 20:01:08", "apply recorded adjustments", "10", "10", "COMPLETE", "thermal + alignment fields written"],
    ["06", "2027-03-16 20:01:09", "send to temperature-correction consumer", "10", "10", "COMPLETE", "pipeline run accepted"],
]
with open(os.path.join(OUT, "FOREMAN-M13-pipeline-log.csv"), "w", newline="") as f:
    csv.writer(f, lineterminator="\n").writerows(log_rows)

print(
    "  wrote M13 handoff: "
    f"{len(sensor_rows)} sensor rows, {len(weather_rows)} weather rows, "
    f"{len(pipeline_rows)} emitted model rows"
)
